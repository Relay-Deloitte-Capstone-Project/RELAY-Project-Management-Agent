"""Generic PM-document upload -> classify -> chunk -> review -> index.

Covers the 11 non-SOW doc types from the PM Tool Ingestion & Retrieval
reference (Charter, Deliverables Matrix, BRD/PRD, Change Request, Epic
Brief, Sprint Planning/Review, Retro, Status Report, Risk Log, UAT
Sign-off, Meeting Notes). SOW keeps its own existing pipeline (api/sow.py) —
this module never touches sow_documents/sow_deliverables.

Pipeline, per uploaded file:
  1. Extract raw text (.md/.txt directly, .pdf via pypdf, .docx via
     python-docx — headings are reconstructed as `## `-prefixed lines and
     tables as markdown pipe-tables so the same downstream chunker handles
     all three formats identically).
  2. Stage 1 (fast path, free): if the file has valid YAML front matter
     naming a known doc_type, trust it and chunk deterministically by
     markdown structure (_markdown_chunk) + regex entity extraction.
  3. Stage 2 (fallback): no usable front matter -> one structured LLM call
     (api.llm.classify_document) does classification, field extraction,
     section-chunking, and entity extraction together.
  4. Store the result on pm_documents.pending_sections — NOT into
     public.chunks yet. Nothing is searchable/citable until a human
     confirms it via PATCH .../confirm: an unreviewed classification could
     have the wrong doc_type or a garbled section split, and this is the
     same abstain-rather-than-guess principle the query pipeline already
     applies to answering, applied here to ingestion.
  5. On confirm: embed each pending section, insert into public.chunks
     tagged with this doc's entities/section_path, mark is_latest=TRUE, and
     demote any earlier CONFIRMED document in the same recurrence group
     (same engagement_id + doc_type + recurrence_key, or + doc_id) to
     is_latest=FALSE — both the pm_documents row and its chunks, since
     query.py reads is_latest off chunks directly to avoid a join.
"""

import hashlib
import io
import json
import logging
import os
import re
import uuid
import zipfile
from pathlib import Path
from typing import Optional

import asyncpg
import yaml
from fastapi import APIRouter, File, Form, HTTPException, Request, UploadFile
from pydantic import BaseModel
from docx import Document as DocxDocument
from pypdf import PdfReader
from starlette.concurrency import run_in_threadpool

from api import llm
from api.query import embed, ready_model, to_pgvector

logger = logging.getLogger(__name__)
router = APIRouter()

DEFAULT_ENGAGEMENT_ID = os.environ.get("RELAY_ENGAGEMENT_ID", "proj-001")
STORAGE_ROOT = Path(__file__).parent.parent / "storage" / "pm_documents"
MAX_UPLOAD_BYTES = 20 * 1024 * 1024

KNOWN_DOC_TYPES = set(llm.PM_DOC_TYPES)  # includes "unclassified"
ENTITY_KEYS = [
    "people", "deliverable_ids", "requirement_ids", "risk_ids",
    "ticket_refs", "decision_ids", "cr_refs",
]

# Confidence below this still ingests (never silently drops content) but the
# review UI should flag it harder than a routine review — see GET response's
# `needs_attention`.
LOW_CONFIDENCE = 0.6

# Target chunk size for the deterministic markdown chunker, mirroring the
# reference doc's 150-400 *token* guidance in characters (~4 chars/token).
CHUNK_TARGET_CHARS = 1600
CHUNK_OVERLAP = 150

# Doc types whose canonical structure (§3 of the reference) is "one table,
# one row = one chunk" — splitting these by `##` heading alone would still
# leave an entire table as one oversized, un-filterable chunk.
ROW_CHUNKED_TYPES = {"deliverables_matrix", "risk_log"}

FRONT_MATTER_RE = re.compile(r"\A---\s*\n(.*?)\n---\s*\n?", re.DOTALL)
HEADING_RE = re.compile(r"^##\s+(.+)$", re.MULTILINE)
TABLE_ROW_RE = re.compile(r"^\|.*\|\s*$", re.MULTILINE)

ENTITY_PATTERNS = {
    "ticket_refs": re.compile(r"\bKPD-\d+\b"),
    "risk_ids": re.compile(r"\b[RI]-\d+\b"),
    "cr_refs": re.compile(r"\bCR-\d+\b"),
    "decision_ids": re.compile(r"\bD-\d+\b"),
    "requirement_ids": re.compile(r"\b(?:FR|NFR)-\d+\b"),
    "deliverable_ids": re.compile(r"\bD\d+\b(?!-)"),
}


def _friendly_error(exc: Exception) -> str:
    text = str(exc)
    if "RESOURCE_EXHAUSTED" in text or "429" in text:
        return "LLM quota exceeded. Wait for it to reset or switch providers, then retry."
    return text[:300] + "…" if len(text) > 300 else text


def _docx_table_to_markdown(table) -> str:
    rows = [[cell.text.strip().replace("\n", " ") for cell in row.cells] for row in table.rows]
    if not rows:
        return ""
    header, body_rows = rows[0], rows[1:]
    lines = [
        "| " + " | ".join(header) + " |",
        "|" + "|".join(["---"] * len(header)) + "|",
    ]
    lines += ["| " + " | ".join(r) + " |" for r in body_rows]
    return "\n".join(lines)


def _extract_docx(raw: bytes) -> str:
    """Reconstructs a .docx as the same markdown shape .md/.pdf produce —
    Word "Heading N" styles become `## ` lines and tables become pipe-tables
    — so _markdown_chunk / _split_table_rows work identically regardless of
    which format a document arrived in. python-docx exposes paragraphs and
    tables as separate top-level sequences, not interleaved in document
    order, so this walks the underlying XML body to preserve their real
    order — otherwise every table would get shifted to the end of the doc."""
    doc = DocxDocument(io.BytesIO(raw))
    parts = []
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = next((p for p in doc.paragraphs if p._p is child), None)
            if para is None or not para.text.strip():
                continue
            style = (para.style.name or "") if para.style else ""
            if style.lower().startswith("heading"):
                parts.append("## " + para.text.strip())
            else:
                parts.append(para.text.strip())
        elif tag == "tbl":
            table = next((t for t in doc.tables if t._tbl is child), None)
            if table is not None:
                md = _docx_table_to_markdown(table)
                if md:
                    parts.append(md)
    return "\n\n".join(parts)


def _extract_text(filename: str, raw: bytes) -> str:
    lower = filename.lower()
    if lower.endswith((".md", ".markdown", ".txt")):
        try:
            return raw.decode("utf-8")
        except UnicodeDecodeError:
            return raw.decode("latin-1")
    if lower.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(raw))
        pages = [page.extract_text() or "" for page in reader.pages]
        if not any(p.strip() for p in pages):
            raise RuntimeError("No extractable text found in PDF (scanned image, not text-based)")
        return "\n\n".join(pages)
    if lower.endswith(".docx"):
        return _extract_docx(raw)
    raise RuntimeError(
        "Unsupported file type for '{}' — .md, .txt, .pdf, and .docx are supported".format(filename)
    )


def _parse_front_matter(text: str):
    """Returns (metadata_dict, body) if the file opens with a valid YAML
    front-matter block naming a known doc_type; otherwise (None, text)."""
    m = FRONT_MATTER_RE.match(text)
    if not m:
        return None, text
    try:
        meta = yaml.safe_load(m.group(1))
    except yaml.YAMLError:
        return None, text
    if not isinstance(meta, dict) or meta.get("doc_type") not in KNOWN_DOC_TYPES:
        return None, text
    return meta, text[m.end():]


def _split_table_rows(block: str) -> list:
    """Splits a markdown table into one chunk per row (header + separator
    row excluded) — the reference doc's "one row = one chunk" rule for
    Deliverables Matrix / Risk Log tables, so a status change to one row
    never touches the others' retrieval context. Returns
    [{"row_id": <first cell>, "content": <header + row>}, ...]."""
    lines = [ln for ln in block.splitlines() if ln.strip()]
    rows = [ln for ln in lines if TABLE_ROW_RE.match(ln)]
    if len(rows) < 3:  # header + separator + at least one data row
        return [{"row_id": None, "content": block}] if block.strip() else []
    header, sep = rows[0], rows[1]
    out = []
    for row in rows[2:]:
        cells = [c.strip() for c in row.strip().strip("|").split("|")]
        row_id = cells[0] if cells and cells[0] else None
        out.append({"row_id": row_id, "content": "\n".join([header, sep, row])})
    return out


def _split_long_section(text: str) -> list:
    """Fallback sliding-window split for a section that's still too long
    after heading/row splitting — same overlap approach as sow.py's
    _chunk_page, applied only to the rare oversized leftover instead of the
    whole document, so most chunks stay heading/row-aligned."""
    text = text.strip()
    if len(text) <= CHUNK_TARGET_CHARS:
        return [text] if text else []
    pieces = []
    start = 0
    while start < len(text):
        end = min(start + CHUNK_TARGET_CHARS, len(text))
        piece = text[start:end].strip()
        if piece:
            pieces.append(piece)
        if end == len(text):
            break
        start = end - CHUNK_OVERLAP
    return pieces


def _markdown_chunk(body: str, doc_type: str) -> list:
    """Deterministic Stage-1 chunker for a file with valid front matter:
    split on `##` headings, then row-split tables for the doc types where a
    table row *is* the record, then sliding-window anything still oversized.
    Returns [{"section_path": ..., "content": ...}, ...]."""
    heading_positions = [(m.start(), m.group(1).strip()) for m in HEADING_RE.finditer(body)]
    if not heading_positions:
        sections = [("(document)", body)]
    else:
        sections = []
        # Anything before the first heading (a title line, an intro sentence).
        if heading_positions[0][0] > 0:
            preamble = body[: heading_positions[0][0]].strip()
            if preamble:
                sections.append(("(preamble)", preamble))
        for i, (pos, heading) in enumerate(heading_positions):
            end = heading_positions[i + 1][0] if i + 1 < len(heading_positions) else len(body)
            sections.append((heading, body[pos:end]))

    chunks = []
    for heading, content in sections:
        content = content.strip()
        if not content:
            continue
        if doc_type in ROW_CHUNKED_TYPES and TABLE_ROW_RE.search(content):
            doc_label = doc_type.replace("_", " ").title()
            for row in _split_table_rows(content):
                path = "{} > {}".format(doc_label, row["row_id"]) if row["row_id"] else heading
                chunks.append({"section_path": path, "content": row["content"]})
        else:
            for piece in _split_long_section(content):
                chunks.append({"section_path": heading, "content": piece})
    return chunks


def _regex_entities(text: str) -> dict:
    return {
        key: sorted(set(pattern.findall(text)))
        for key, pattern in ENTITY_PATTERNS.items()
    } | {"people": []}  # names need real NER; not guessed via regex


async def _log(pool: asyncpg.Pool, engagement_id: str, source: str, message: str, level: str = "info"):
    await pool.execute(
        "INSERT INTO public.ingestion_logs (engagement_id, source, message, level) VALUES ($1, $2, $3, $4)",
        engagement_id, source, message, level,
    )


def _row_to_dict(r) -> dict:
    d = dict(r)
    for k in ("id", "pm_document_id"):
        if k in d and d[k] is not None:
            d[k] = str(d[k])
    for k in ("doc_date", "uploaded_at", "confirmed_at"):
        if d.get(k) is not None:
            d[k] = d[k].isoformat()
    for k in ("entities", "pending_sections"):
        if isinstance(d.get(k), str):
            d[k] = json.loads(d[k])
    return d


DOC_SELECT = """
    SELECT id, engagement_id, doc_type, doc_id, scenario, doc_version, doc_date, author,
           doc_status, recurrence_key, is_latest, source_file_name, ingestion_status,
           classification_confidence, classification_notes, parse_error, entities,
           pending_sections, uploaded_by, uploaded_at, confirmed_by, confirmed_at
    FROM public.pm_documents
"""


async def _ingest_one(
    pool: asyncpg.Pool, engagement_id: str, filename: str, raw: bytes, uploaded_by: Optional[str]
) -> dict:
    document_id = uuid.uuid4()
    file_sha256 = hashlib.sha256(raw).hexdigest()
    ext = "".join(Path(filename).suffixes)[-10:] or ".txt"
    doc_dir = STORAGE_ROOT / engagement_id
    doc_dir.mkdir(parents=True, exist_ok=True)
    storage_path = doc_dir / "{}{}".format(document_id, ext)
    storage_path.write_bytes(raw)

    row = await pool.fetchrow(
        """
        INSERT INTO public.pm_documents
            (id, engagement_id, doc_type, source_file_name, storage_path, file_sha256,
             ingestion_status, uploaded_by)
        VALUES ($1, $2, 'unclassified', $3, $4, $5, 'classifying', $6)
        RETURNING id
        """,
        document_id, engagement_id, filename, str(storage_path), file_sha256, uploaded_by,
    )

    try:
        text = _extract_text(filename, raw)
        if not text.strip():
            raise RuntimeError("File is empty or has no extractable text")

        front_matter, body = _parse_front_matter(text)
        if front_matter is not None:
            doc_type = front_matter["doc_type"]
            sections = _markdown_chunk(body, doc_type)
            entities = _regex_entities(body)
            confidence = 1.0
            notes = "Front matter present and valid — deterministic parse, no LLM call."
            fields = {
                "doc_id": front_matter.get("doc_id"),
                "scenario": front_matter.get("scenario"),
                "doc_version": str(front_matter.get("version") or front_matter.get("doc_version") or "") or None,
                "doc_date": front_matter.get("date") or front_matter.get("doc_date"),
                "author": front_matter.get("author"),
                "doc_status": front_matter.get("status") or front_matter.get("doc_status"),
                "recurrence_key": (
                    str(front_matter.get("week_start") or front_matter.get("sprint_number") or "") or None
                ),
            }
        else:
            result = await llm.classify_document(text[:60000])
            doc_type = result.get("doc_type") or "unclassified"
            if doc_type not in KNOWN_DOC_TYPES:
                doc_type = "unclassified"
            # For table-structured types, row-splitting is mechanical and
            # unambiguous — the LLM has no judgment call to make that a regex
            # can't make more cheaply and more auditably, and getting a
            # financial/risk row's split wrong is exactly where "trust the
            # model" is least acceptable. Prefer the deterministic splitter
            # whenever the raw text actually contains a markdown table;
            # otherwise (e.g. a matrix described in prose, no real table)
            # fall back to the LLM's own section split — there's no
            # structure for the deterministic path to find in that case.
            if doc_type in ROW_CHUNKED_TYPES and TABLE_ROW_RE.search(text):
                sections = _markdown_chunk(text, doc_type)
            else:
                sections = result.get("sections") or []
            entities = {k: (result.get("entities") or {}).get(k) or [] for k in ENTITY_KEYS}
            confidence = float(result.get("confidence") or 0.0)
            # Never empty — the frontend falls back to a hardcoded "via
            # front matter" message when this is falsy, which would be
            # actively wrong here (this is the LLM path, front matter was
            # absent or invalid). An empty string from the model still needs
            # its own honest default.
            notes = result.get("notes") or "Classified via one LLM call — the model gave no extra reasoning."
            fields = {
                "doc_id": result.get("doc_id"),
                "scenario": result.get("scenario"),
                "doc_version": result.get("doc_version"),
                "doc_date": result.get("doc_date") or None,
                "author": result.get("author"),
                "doc_status": result.get("doc_status"),
                "recurrence_key": result.get("recurrence_key"),
            }

        # Both paths can hand back "" for an unset field (YAML front matter
        # with a blank value; the LLM occasionally emitting "" rather than
        # omitting a key despite the prompt) — doc_status and scenario are
        # CHECK-constrained columns that reject "" outright (only NULL or a
        # real enum value), so this must run before the INSERT, not after.
        fields = {k: (v or None) for k, v in fields.items()}
        if fields["scenario"] not in (None, "client", "internal"):
            fields["scenario"] = None
        if fields["doc_status"] not in (None, "draft", "final", "superseded"):
            fields["doc_status"] = None

        if not sections:
            raise RuntimeError("Parsed but produced no content sections")

        # doc_date arrives as either a native date (YAML parses an unquoted
        # `date: 2026-09-14` front-matter value into one already) or free
        # text from the LLM path — normalize both to a real date object
        # (asyncpg's DATE codec rejects a plain str outright) and only trust
        # it if it's valid, otherwise leave it null rather than fail the
        # whole ingestion over a cosmetic field.
        from datetime import date as _date
        raw_doc_date = fields["doc_date"]
        doc_date = None
        if isinstance(raw_doc_date, _date):
            doc_date = raw_doc_date
        elif raw_doc_date:
            try:
                doc_date = _date.fromisoformat(str(raw_doc_date))
            except ValueError:
                doc_date = None

        await pool.execute(
            """
            UPDATE public.pm_documents SET
                doc_type = $2, doc_id = $3, scenario = $4, doc_version = $5, doc_date = $6,
                author = $7, doc_status = $8, recurrence_key = $9,
                ingestion_status = 'needs_review', classification_confidence = $10,
                classification_notes = $11, entities = $12::jsonb, pending_sections = $13::jsonb
            WHERE id = $1
            """,
            document_id, doc_type, fields["doc_id"], fields["scenario"], fields["doc_version"],
            doc_date, fields["author"], fields["doc_status"], fields["recurrence_key"],
            confidence, notes, json.dumps(entities), json.dumps(sections),
        )
        await _log(
            pool, engagement_id, "pm_documents",
            "Ingested {} as {} ({} sections, confidence {:.2f})".format(
                filename, doc_type, len(sections), confidence
            ),
        )
    except Exception as exc:
        await pool.execute(
            "UPDATE public.pm_documents SET ingestion_status = 'failed', parse_error = $2 WHERE id = $1",
            document_id, _friendly_error(exc),
        )
        await _log(pool, engagement_id, "pm_documents", "Failed on {}: {}".format(filename, exc), level="error")

    result_row = await pool.fetchrow(DOC_SELECT + " WHERE id = $1", document_id)
    return _row_to_dict(result_row)


def _iter_upload_files(filename: str, raw: bytes):
    """Yields (name, bytes) — expands a .zip into its member files (a
    'compressed folder' upload) so the pipeline below never needs to know
    whether the request was individual files or a zipped folder."""
    if filename.lower().endswith(".zip"):
        with zipfile.ZipFile(io.BytesIO(raw)) as zf:
            for info in zf.infolist():
                name = info.filename
                if info.is_dir() or name.startswith("__MACOSX/") or Path(name).name.startswith("."):
                    continue
                yield Path(name).name, zf.read(info)
    else:
        yield filename, raw


@router.post("/api/admin/documents/upload")
async def upload_documents(
    request: Request,
    files: list[UploadFile] = File(...),
    engagement_id: str = Form(DEFAULT_ENGAGEMENT_ID),
    uploaded_by: Optional[str] = Form(None),
):
    """Accepts one or more files in a single request — either picked
    individually, picked as a folder (browsers send a folder's files as a
    flat multi-file list), or a single .zip of a folder. Each resulting file
    is classified independently; one bad file in a batch doesn't fail the
    others (each gets its own try/except in _ingest_one)."""
    pool: asyncpg.Pool = request.app.state.pool
    project = await pool.fetchval(
        "SELECT engagement_id FROM public.projects WHERE engagement_id = $1", engagement_id
    )
    if project is None:
        raise HTTPException(status_code=404, detail="Unknown engagement_id — no such project")

    results = []
    for upload in files:
        raw = await upload.read()
        if len(raw) > MAX_UPLOAD_BYTES:
            results.append({"source_file_name": upload.filename, "error": "File exceeds 20MB limit"})
            continue
        if not raw:
            continue
        for name, member_bytes in _iter_upload_files(upload.filename or "upload", raw):
            if not member_bytes:
                continue
            results.append(await _ingest_one(pool, engagement_id, name, member_bytes, uploaded_by))

    return {"ingested": results}


@router.get("/api/admin/documents")
async def list_documents(
    request: Request,
    engagement_id: str = DEFAULT_ENGAGEMENT_ID,
    ingestion_status: Optional[str] = None,
    doc_type: Optional[str] = None,
):
    pool: asyncpg.Pool = request.app.state.pool
    query = DOC_SELECT + " WHERE engagement_id = $1"
    params = [engagement_id]
    if ingestion_status:
        params.append(ingestion_status)
        query += " AND ingestion_status = ${}".format(len(params))
    if doc_type:
        params.append(doc_type)
        query += " AND doc_type = ${}".format(len(params))
    query += " ORDER BY uploaded_at DESC"
    rows = await pool.fetch(query, *params)
    docs = [_row_to_dict(r) for r in rows]
    for d in docs:
        d["needs_attention"] = (
            d["ingestion_status"] == "needs_review"
            and (d["classification_confidence"] is None or d["classification_confidence"] < LOW_CONFIDENCE
                 or d["doc_type"] == "unclassified")
        )
    return {"documents": docs}


@router.get("/api/admin/documents/{document_id}")
async def get_document(document_id: str, request: Request):
    pool: asyncpg.Pool = request.app.state.pool
    row = await pool.fetchrow(DOC_SELECT + " WHERE id = $1", document_id)
    if row is None:
        raise HTTPException(status_code=404, detail="Document not found")
    return _row_to_dict(row)


class ConfirmEdits(BaseModel):
    doc_type: Optional[str] = None
    doc_id: Optional[str] = None
    scenario: Optional[str] = None
    doc_version: Optional[str] = None
    doc_date: Optional[str] = None
    author: Optional[str] = None
    doc_status: Optional[str] = None
    recurrence_key: Optional[str] = None
    pending_sections: Optional[list] = None
    confirmed_by: Optional[str] = None


@router.patch("/api/admin/documents/{document_id}/confirm")
async def confirm_document(document_id: str, body: ConfirmEdits, request: Request):
    """Applies any last edits, then materializes pending_sections into
    public.chunks (embedding each one) and flips is_latest for this
    recurrence group — this is the only place a PM document becomes
    citable in Ask Project."""
    pool: asyncpg.Pool = request.app.state.pool
    embedding_model = await ready_model(request)

    row = await pool.fetchrow(DOC_SELECT + " WHERE id = $1", document_id)
    if row is None:
        raise HTTPException(status_code=404, detail="Document not found")
    doc = _row_to_dict(row)
    if doc["ingestion_status"] not in ("needs_review", "confirmed"):
        raise HTTPException(
            status_code=409,
            detail="Document is '{}' — only needs_review documents can be confirmed".format(
                doc["ingestion_status"]
            ),
        )

    doc_type = body.doc_type or doc["doc_type"]
    if doc_type not in KNOWN_DOC_TYPES or doc_type == "unclassified":
        raise HTTPException(status_code=400, detail="Set a real doc_type before confirming")
    sections = body.pending_sections if body.pending_sections is not None else doc["pending_sections"]

    # A reviewer correcting a misclassification (e.g. the LLM said
    # "unclassified" for a sparse Deliverables Matrix, and a human fixes it
    # to the real type here) wouldn't otherwise get the deterministic
    # row-per-chunk splitting that _ingest_one applies for row-chunked types
    # at upload time — the sections already on the document reflect whatever
    # the original classification produced. Re-derive from the stored file
    # whenever the resolved type is row-chunked and the reviewer didn't
    # submit their own hand-edited sections (an explicit edit always wins).
    if doc_type in ROW_CHUNKED_TYPES and body.pending_sections is None:
        storage_path = await pool.fetchval(
            "SELECT storage_path FROM public.pm_documents WHERE id = $1", document_id
        )
        try:
            raw = Path(storage_path).read_bytes()
            text = _extract_text(doc["source_file_name"], raw)
            if TABLE_ROW_RE.search(text):
                sections = _markdown_chunk(text, doc_type)
        except Exception:
            logger.warning("Could not re-derive row chunks for %s on confirm", document_id, exc_info=True)

    if not sections:
        raise HTTPException(status_code=400, detail="No sections to index")

    recurrence_key = body.recurrence_key if body.recurrence_key is not None else doc["recurrence_key"]
    doc_id_field = body.doc_id if body.doc_id is not None else doc["doc_id"]

    # A document with neither a doc_id nor a recurrence_key (e.g. generic
    # meeting notes with no formal ID and no extracted week/sprint number)
    # would never match its own group below — every upload would stay
    # is_latest=TRUE forever instead of exactly one per group. Fall back to
    # doc_date, then the source filename, so there's always *some* grouping
    # key: two uploads of the literal same meeting notes (same date/file)
    # correctly supersede, while two different real meetings (different
    # dates) correctly don't.
    if not recurrence_key and not doc_id_field:
        recurrence_key = (
            (body.doc_date or doc["doc_date"]) if (body.doc_date or doc["doc_date"]) else doc["source_file_name"]
        )

    async with pool.acquire() as conn:
        async with conn.transaction():
            # Demote the previous latest in this group BEFORE inserting the
            # new chunks, so there's never a window with two is_latest rows.
            group_filter = "doc_type = $2 AND (recurrence_key = $3 OR (recurrence_key IS NULL AND $3 IS NULL AND doc_id = $4))"
            old_docs = await conn.fetch(
                "SELECT id FROM public.pm_documents WHERE engagement_id = $1 AND {} "
                "AND ingestion_status = 'confirmed' AND id != $5".format(group_filter),
                doc["engagement_id"], doc_type, recurrence_key, doc_id_field, document_id,
            )
            if old_docs:
                old_ids = [r["id"] for r in old_docs]
                await conn.execute(
                    "UPDATE public.pm_documents SET is_latest = FALSE WHERE id = ANY($1::uuid[])", old_ids
                )
                await conn.execute(
                    "UPDATE public.chunks SET is_latest = FALSE WHERE pm_document_id = ANY($1::uuid[])", old_ids
                )

            # Remove any chunks from a prior confirm of *this* document
            # (re-confirming after an edit shouldn't duplicate them).
            await conn.execute("DELETE FROM public.chunks WHERE pm_document_id = $1", document_id)

            for section in sections:
                content = (section.get("content") or "").strip()
                if not content:
                    continue
                section_path = section.get("section_path")
                # A row-chunk's section_path is "<Doc Type> > <row id>" (see
                # _markdown_chunk / the classify_document prompt) — use that
                # row id as this chunk's own source_doc_id so it's citable as
                # [R-01], not the whole document's id. Every row of a Risk
                # Log otherwise shares one source_doc_id and the model's
                # per-row citation (which is what ENTITY_LOOKUP_SQL in
                # intents.py depends on to answer "status of R-01") would
                # never match any chunk's citable id.
                row_id = section_path.rsplit(" > ", 1)[-1] if section_path and " > " in section_path else None
                source_doc_id = row_id or doc_id_field or str(document_id)
                # Per-CHUNK entities, not the whole document's — copying the
                # document-level entity summary onto every chunk meant a Risk
                # Log row that merely *mentions* "CR-002" in its mitigation
                # text got tagged as if it were CR-002's own record, and
                # ENTITY_LOOKUP_SQL (intents.py) had no way to tell a genuine
                # match from an incidental one. Re-deriving entities from
                # just this chunk's own text keeps the exact-ID lookup
                # precise; doc["entities"] (the document-level union) stays
                # on pm_documents for the review UI's summary view only.
                chunk_entities = _regex_entities(content)
                vec = await run_in_threadpool(embed, embedding_model, content)
                await conn.execute(
                    """
                    INSERT INTO public.chunks
                        (engagement_id, source_type, source_doc_id, content, metadata,
                         embedding, pm_document_id, is_latest, section_path, entities)
                    VALUES ($1, $2, $3, $4, $5::jsonb, $6::vector, $7, TRUE, $8, $9::jsonb)
                    """,
                    doc["engagement_id"], doc_type, source_doc_id, content,
                    json.dumps({"pm_document_id": str(document_id)}), to_pgvector(vec),
                    document_id, section_path, json.dumps(chunk_entities),
                )

            await conn.execute(
                """
                UPDATE public.pm_documents SET
                    doc_type = $2, doc_id = $3, scenario = COALESCE($4, scenario),
                    doc_version = COALESCE($5, doc_version), author = COALESCE($6, author),
                    doc_status = COALESCE($7, doc_status), recurrence_key = $8,
                    pending_sections = $9::jsonb, ingestion_status = 'confirmed',
                    is_latest = TRUE, confirmed_by = $10, confirmed_at = NOW()
                WHERE id = $1
                """,
                document_id, doc_type, doc_id_field, body.scenario, body.doc_version,
                body.author, body.doc_status, recurrence_key, json.dumps(sections),
                body.confirmed_by,
            )

    await _log(
        pool, doc["engagement_id"], "pm_documents",
        "Confirmed {} as {} ({} chunks indexed)".format(doc["source_file_name"], doc_type, len(sections)),
    )
    result_row = await pool.fetchrow(DOC_SELECT + " WHERE id = $1", document_id)
    return _row_to_dict(result_row)


@router.patch("/api/admin/documents/{document_id}/reject")
async def reject_document(document_id: str, request: Request):
    pool: asyncpg.Pool = request.app.state.pool
    result = await pool.execute(
        "UPDATE public.pm_documents SET ingestion_status = 'rejected' WHERE id = $1", document_id
    )
    if result == "UPDATE 0":
        raise HTTPException(status_code=404, detail="Document not found")
    return {"ok": True}


@router.delete("/api/admin/documents/{document_id}")
async def delete_document(document_id: str, request: Request):
    pool: asyncpg.Pool = request.app.state.pool
    result = await pool.execute("DELETE FROM public.pm_documents WHERE id = $1", document_id)
    if result == "DELETE 0":
        raise HTTPException(status_code=404, detail="Document not found")
    return {"ok": True}
